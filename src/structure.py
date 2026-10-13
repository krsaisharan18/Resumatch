"""
Layout- and structure-aware helpers for ResuMatch.

Why this module exists
----------------------
Resume PDFs come in wildly different layouts (sidebars, timelines with a date
column, two-column pages, skill "chips", tables). Reading them as one long
stream of text interleaves columns and loses which date belongs to which job.
This module fixes that in three steps, with no ML and no LLM:

1. ``pdf_text`` / ``docx_text``  - read the file in reading order (columns are
   detected from word positions, tables are kept row by row).
2. ``prepare_text``               - clean extraction artefacts such as "(cid:127)".
3. ``parse_*``                    - turn each section into structured entries
   (title / company / dates / location / bullets, degree / institution / years,
   certificate / issuer / date) using date lines as anchors.

GLiNER (when installed) still finds names, organisations and locations; the
structured result is used to validate it and to fill anything it missed.
"""
import re
from statistics import median

# ══════════════════════════════════════════════════════════════════════════════
# 1. Reading files in reading order
# ══════════════════════════════════════════════════════════════════════════════

_GAP = "  "          # marker for a wide horizontal gap inside one line (chips, columns)


def _group_lines(words, tol=3.0):
    """Group pdfplumber words into visual lines (top-to-bottom, left-to-right).
    Words share a line when they share a baseline (bottom), so a name set in small
    caps / big first letters ('A' 30pt + 'MAN' 22pt) stays on one line."""
    words = sorted(words, key=lambda w: (round(w["bottom"]), w["x0"]))
    lines, cur, cur_b = [], [], None
    for w in words:
        if cur_b is None or abs(w["bottom"] - cur_b) <= tol:
            cur.append(w)
            cur_b = w["bottom"] if cur_b is None else (cur_b + w["bottom"]) / 2
        else:
            lines.append(sorted(cur, key=lambda x: x["x0"]))
            cur, cur_b = [w], w["bottom"]
    if cur:
        lines.append(sorted(cur, key=lambda x: x["x0"]))
    return lines


def _line_text(ws):
    """Join the words of one line, marking wide gaps with a double space."""
    if len(ws) >= 5 and sum(len(w["text"]) == 1 for w in ws) / len(ws) >= 0.8:
        # letter-spaced heading: "B E L L A   T R E V I N O" -> "BELLA TREVINO"
        gaps = [b["x0"] - a["x1"] for a, b in zip(ws, ws[1:])]
        g = median(gaps)
        out = ws[0]["text"]
        for gap, b in zip(gaps, ws[1:]):
            out += (" " if gap > g * 1.8 + 0.5 else "") + b["text"]
        return out
    toks = [w["text"] for w in ws]
    if 3 <= len(ws) <= 8 and all(t.isalpha() for t in toks) and (
            any(len(t) == 1 for t in toks) or any(re.search(r"[A-Z]{2,}[a-z]", t) for t in toks)):
        # name set in small caps / big first letters: 'A MAN S HARM A' -> 'AMAN SHARMA'.
        # Fragments of one word sit close together; real word gaps are much wider.
        gaps = [b["x0"] - a["x1"] for a, b in zip(ws, ws[1:])]
        mx = max(gaps)
        size_ref = max(w.get("size", 9) for w in ws)
        if mx > 0:
            out = ws[0]["text"]
            for gap, b in zip(gaps, ws[1:]):
                out += (" " if gap >= 0.35 * mx or gap > 1.5 and gap >= size_ref * 0.28 else "") + b["text"]
            if " " in out:
                return out
    out = ws[0]["text"]
    for a, b in zip(ws, ws[1:]):
        size = max(a.get("size", 9), b.get("size", 9))
        gap = b["x0"] - a["x1"]
        out += (_GAP if gap > size * 0.6 else " ") + b["text"]
    return out


def _find_gutter(lines, page_w):
    """x position of a vertical gutter that splits the page into two real
    columns (main + sidebar, or two equal columns), or None. Lines that run
    across the gutter (banner, summary, footer) are ignored when judging, so a
    full-width header does not hide a two-column body. A narrow left column is
    accepted only when its lines are mostly NOT on the same baseline as the
    right column (a sidebar), never for a date column that sits beside each
    entry (a timeline)."""
    words = [w for ln in lines for w in ln]
    if len(words) < 40:
        return None
    best = None
    for pct in range(18, 71):
        x = page_w * pct / 100
        left = right = crossing = 0
        left_lines = right_lines = paired = 0
        starts = []
        for ln in lines:
            l = [w for w in ln if w["x1"] <= x + 1]
            r = [w for w in ln if w["x0"] >= x - 1]
            straddle = any(w["x0"] < x - 1 and w["x1"] > x + 1 for w in ln)
            tight = bool(l and r) and min(w["x0"] for w in r) - max(w["x1"] for w in l) < 14
            if straddle or tight:
                crossing += 1
                continue
            left += len(l); right += len(r)
            left_lines += bool(l); right_lines += bool(r)
            paired += bool(l and r)
            if r:
                starts.append(round(min(w["x0"] for w in r) / 4))
        total = left + right
        if total == 0 or crossing > 0.35 * len(lines) or not starts:
            continue
        share = min(left, right) / total
        if min(left_lines, right_lines) < 8 or share < 0.1:
            continue
        if share < 0.28 and paired > 0.5 * min(left_lines, right_lines):
            continue
        common = max(starts.count(v) for v in set(starts)) / len(starts)
        if common < 0.45 or len(starts) < max(8, 0.3 * len(lines)):
            continue
        score = share - 0.5 * crossing / max(len(lines), 1)
        if best is None or score > best[0] + 1e-9:
            best = (score, x)
    return best[1] if best else None


def _split_columns(lines, x):
    """Header lines first, then the whole left column, then the whole right column.

    A line belongs to the header when it is continuous text across the gutter
    (summary, name, contact band). The body starts at the first line whose right
    part begins at the right column's usual left edge."""
    def crosses(ln):
        return any(w["x0"] < x - 1 and w["x1"] > x + 1 for w in ln)

    def right_start(ln):
        r = [w for w in ln if w["x0"] >= x - 1]
        return min(w["x0"] for w in r) if r else None

    def real_gap(ln):
        l = [w for w in ln if w["x1"] <= x + 1]
        r = [w for w in ln if w["x0"] >= x - 1]
        if not r or crosses(ln):
            return False
        return (not l) or min(w["x0"] for w in r) - max(w["x1"] for w in l) >= 14

    starts = [round(right_start(ln) / 4) for ln in lines if real_gap(ln)]
    if not starts:
        return lines
    cnt = {v: starts.count(v) for v in set(starts)}
    mode = min(v for v, c in cnt.items() if c >= max(3, 0.2 * max(cnt.values())))   # the column's left margin
    first_body = next((i for i, ln in enumerate(lines)
                       if real_gap(ln) and abs(round(right_start(ln) / 4) - mode) <= 1), None)
    if first_body is None:
        return lines
    head, body = lines[:first_body], lines[first_body:]
    left, right = [], []
    for ln in body:
        if crosses(ln) or (not real_gap(ln) and any(w["x0"] >= x - 1 for w in ln)):
            left.append(ln)          # a full-width row inside the body stays with the left column
            continue
        l = [w for w in ln if w["x1"] <= x + 1]
        r = [w for w in ln if w["x0"] >= x - 1]
        if l: left.append(l)
        if r: right.append(r)
    return head + left + right


def _split_runs_of_two_blocks(lines):
    """Side-by-side text blocks inside a one-column page (e.g. two 'achievement'
    cards): a run of >=3 lines that each have exactly one wide gap at the same
    x, with real sentences on both sides, is read left block then right block.
    Date/location gutters (short text on the left) are left untouched."""
    out, i, n = [], 0, len(lines)

    def gap_of(ln):
        gaps = [(b["x0"] - a["x1"], a["x1"], b["x0"]) for a, b in zip(ln, ln[1:])]
        big = [g for g in gaps if g[0] > 28]
        if len(big) != 1:
            return None
        g = big[0]
        return g[2], ln.index(next(w for w in ln if w["x0"] >= g[2] - 0.01))

    while i < n:
        g = gap_of(lines[i])
        if g is None:
            out.append(lines[i]); i += 1; continue
        run = [i]
        j = i + 1
        while j < n and (g2 := gap_of(lines[j])) is not None and abs(g2[0] - g[0]) < 8:
            run.append(j); j += 1
        if len(run) >= 3:
            lw = [gap_of(lines[k])[1] for k in run]
            rw = [len(lines[k]) - gap_of(lines[k])[1] for k in run]
            if median(lw) >= 6 and median(rw) >= 6 and min(lw) >= 3 and min(rw) >= 3:
                L = [lines[k][:gap_of(lines[k])[1]] for k in run]
                R = [lines[k][gap_of(lines[k])[1]:] for k in run]
                out.extend(L + R)
                i = j
                continue
        out.extend(lines[k] for k in run)
        i = j
    return out



# ══════════════════════════════════════════════════════════════════════════════
# Scanned / image-only PDFs  ->  OCR (RapidOCR, with pytesseract as a fallback)
# ══════════════════════════════════════════════════════════════════════════════

_OCR_ENGINE = None
_OCR_FAILED = False
OCR_ERROR = ""          # why OCR is unavailable (shown in the app instead of a vague message)
OCR_DPI = 200


def _get_ocr():
    global _OCR_ENGINE, _OCR_FAILED, OCR_ERROR
    if _OCR_ENGINE is not None or _OCR_FAILED:
        return _OCR_ENGINE
    try:
        from rapidocr_onnxruntime import RapidOCR
        _OCR_ENGINE = RapidOCR()
    except Exception as e:
        _OCR_FAILED = True
        OCR_ERROR = f"{type(e).__name__}: {e}"[:300]
    return _OCR_ENGINE


_WORDS = None
_SHORT_OK = {"a", "i", "of", "by", "in", "to", "on", "at", "or", "an", "as", "is", "it", "my",
             "and", "for", "the", "with"}


def _load_words():
    """126k English words ordered by frequency (src/data/words_en.txt.gz)."""
    global _WORDS
    if _WORDS is None:
        import gzip, math
        from pathlib import Path
        _WORDS = {}
        try:
            p = Path(__file__).with_name("data") / "words_en.txt.gz"
            words = gzip.open(p, "rt", encoding="utf-8").read().split()
            n = len(words)
            ln = math.log(n)
            _WORDS = {w: math.log((i + 1) * ln) for i, w in enumerate(words)}
            _WORDS["__max__"] = math.log(30000 * ln)
        except Exception:
            _WORDS = {}
    return _WORDS


def segment_words(token):
    """Split a run-together lowercase string into dictionary words ('thisproject' ->
    ['this','project']). Returns None unless EVERY piece is a real word, so names and
    unknown terms ('kashipur') are never chopped up."""
    W = _load_words()
    if not W or not token.isalpha():
        return None
    t = token.lower()
    if t == "__max__" or t in W or len(t) > 60:
        return None
    n = len(t)
    best = [(0.0, 0)] + [(float("inf"), 0)] * n
    for i in range(1, n + 1):
        for k in range(max(0, i - 20), i):
            piece = t[k:i]
            c = W.get(piece)
            if c is None or c > W["__max__"] or (len(piece) < 3 and piece not in _SHORT_OK):
                continue
            cost = best[k][0] + c
            if cost < best[i][0]:
                best[i] = (cost, k)
    if best[n][0] == float("inf"):
        return None
    out, i = [], n
    while i > 0:
        k = best[i][1]
        out.append(token[k:i]); i = k
    out.reverse()
    return out if len(out) > 1 else None


_OCR_MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september",
               "october", "november", "december"]
_OCR_CONF = str.maketrans({"0": "o", "5": "s", "1": "l", "|": "l", "8": "b", "6": "g"})


def _fix_month_token(tok):
    """'Jon' -> 'Jan', '0ct' -> 'Oct', '5ep' -> 'Sep' (only called right before a year)."""
    import difflib
    t = tok.replace(" ", "").translate(_OCR_CONF).lower()
    if not t or len(t) > 9:
        return None
    for m in _OCR_MONTHS:
        if m == t or m[:3] == t or (m.startswith(t) and len(t) >= 3):
            return m[:3].title() if t == m[:3] or len(t) <= 4 else m.title()
    cands = [m[:3] for m in _OCR_MONTHS] + ["sept"]
    pref = {"o": "a", "a": "o"}
    close = difflib.get_close_matches(t, cands, n=3, cutoff=0.6)
    if not close:
        return None
    if len(close) > 1 and len(t) == 3:                      # tie: prefer the a<->o confusion
        for c in close:
            if all(x == y or pref.get(x) == y for x, y in zip(t, c)):
                return c.title()
    return close[0].title()


def _fix_ocr_dates(out):
    def rep(m):
        fixed = _fix_month_token(m.group(1))
        return f"{fixed} {m.group(2)}" if fixed else m.group(0)
    return re.sub(r"(?<![A-Za-z0-9])([A-Za-z0-9]\s?[A-Za-z]{2,3}|[A-Za-z0-9]{3,9})\s*((?:19|20)\d{2})(?!\d)", rep, out)


def _split_concat(tok, known):
    """Repair run-together OCR tokens, e.g. 'GPKashipur' -> 'GP Kashipur',
    'Courseby' -> 'Course by', 'Leavesofemployee' -> 'Leaves of employee'."""
    # 0) ordinals glued to words: 1stpositionin -> 1st position in
    mo = re.fullmatch(r"(?i)(\d{1,2}(?:st|nd|rd|th))([a-z]{4,})", tok)
    if mo:
        rest = _split_concat(mo.group(2), known)
        return mo.group(1).lower() + " " + rest
    # 0b) ACRONYM + lowercase word: HRMSand -> HRMS and
    mo = re.fullmatch(r"([A-Z]{2,})([a-z]{3,})", tok)
    if mo and mo.group(2) in _load_words():
        return mo.group(1) + " " + mo.group(2)
    # 1) case boundaries and letter/digit boundaries
    t = re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])|(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z])", " ", tok)
    pieces = []
    for p in t.split(" "):
        # 2) HRMSand -> HRMS and
        m = re.fullmatch(r"([A-Z]{2,})([a-z]{3,})", p)
        if m and m.group(2) in _load_words():
            pieces.extend([m.group(1), m.group(2)]); continue
        # 3) dictionary segmentation of lowercase runs
        seg = segment_words(p) if p.isalpha() and len(p) >= 6 else None
        if seg:
            if p[0].isupper():
                seg[0] = seg[0][:1].upper() + seg[0][1:]
            pieces.extend(seg)
        else:
            pieces.append(p)
    # 4) glue pieces that together are a known term (Share+Point -> SharePoint, Word+Press -> WordPress)
    i, merged = 0, []
    while i < len(pieces):
        for n in (3, 2):
            chunk = pieces[i:i + n]
            if len(chunk) == n and "".join(chunk).lower() in known:
                merged.append("".join(chunk)); i += n; break
        else:
            merged.append(pieces[i]); i += 1
    return " ".join(merged)


def _fix_ocr_spacing(text):
    try:
        from src.resume_parser import SKILLS_DB
        known = {k.lower() for k in SKILLS_DB}
    except Exception:
        known = set()
    toks = []
    for tok in text.split(" "):
        core = tok.strip("()[]{}.,;:!?\"'")
        if len(core) >= 7 and not re.search(r"[@/:]|(?i:\.(com|in|net|org|edu|io)\b)", tok):
            lead = tok[:len(tok) - len(tok.lstrip("()[]{}\"'"))]
            trail = tok[len(tok.rstrip(".,;:!?)]}\"'")):]
            tok = lead + _split_concat(core, known) + trail
        toks.append(tok)
    glued = []
    for t in toks:                                  # Share Point / Word Press -> SharePoint / WordPress
        if glued and re.fullmatch(r"[A-Z][a-z]+", glued[-1]) and re.fullmatch(r"[A-Z][a-z]+", t) \
                and (glued[-1] + t).lower() in known:
            glued[-1] += t
        else:
            glued.append(t)
    toks = glued
    out = " ".join(toks)
    out = re.sub(r"\bC\s+#", "C#", out)
    out = _fix_ocr_dates(out)
    out = re.sub(r"(\d{4})\s*-\s*(?=(?i:" + _MON + r"|Present|Pursuing|Current)|\d{4})", r"\1 - ", out)
    out = re.sub(r"(?<=[A-Za-z]{3})\s*-\s*(?=(?:19|20)\d{2}\b)", " - ", out) if re.search(r"(?i)\b" + _MON + r"\s*(?:19|20)\d{2}\s*-", out) else out
    out = re.sub(r"\s+([.,;:!?)])", r"\1", out)
    out = re.sub(r"(?<=\w)\s+(\.)\s*(?=(?:com|in|net|org|co|edu)\b)", ".", out)
    out = re.sub(r"^(?:ist|lst|Ist|1st)\b", "1st", out)
    out = re.sub(r"^[\.\-~·]\s*(?=[A-Z])", "• ", out)
    m = re.fullmatch(r"(?i)([a-z0-9._\-]{3,}?)(?:[o0a@\s]{0,3})(gmail|yahoo|outlook|hotmail|icloud|protonmail|live)\s*\.\s*(com|in|co\.in|net|org)", out.strip())
    if m and "@" not in out:                      # OCR lost the "@"
        local = re.sub(r"(?i)(?:oa|o|a)$", "", m.group(1)) if re.search(r"(?i)[a-z0-9](?:oa|o)$", m.group(1)) else m.group(1)
        out = f"{local}@{m.group(2).lower()}.{m.group(3).lower()}"
    return out


def _tesseract_words(page, dpi=OCR_DPI):
    """Fallback OCR through the tesseract binary (pytesseract), word by word."""
    global OCR_ERROR
    try:
        import pytesseract
        im = page.to_image(resolution=dpi).original.convert("RGB")
        d = pytesseract.image_to_data(im, output_type=pytesseract.Output.DICT)
    except Exception as e:
        OCR_ERROR = (OCR_ERROR + " | " if OCR_ERROR else "") + f"tesseract: {type(e).__name__}: {e}"[:200]
        return []
    k = 72.0 / dpi
    words = []
    for i, txt in enumerate(d["text"]):
        if not txt.strip() or float(d["conf"][i]) < 30:
            continue
        x0, top, w, h = d["left"][i] * k, d["top"][i] * k, d["width"][i] * k, d["height"][i] * k
        words.append({"text": _fix_ocr_spacing(txt.strip()), "x0": x0, "x1": x0 + w, "top": top,
                      "bottom": top + h, "size": max(6.0, h * 0.8), "ocr": True})
    return words


def ocr_words(page, dpi=OCR_DPI):
    """OCR one pdfplumber page -> pseudo-words in PDF points (same shape as extract_words)."""
    eng = _get_ocr()
    if eng is None:
        return _tesseract_words(page, dpi)
    import numpy as np
    im = page.to_image(resolution=dpi).original.convert("RGB")
    res, _ = eng(np.array(im))
    k = 72.0 / dpi
    words = []
    for box, txt, conf in (res or []):
        if not txt or not txt.strip() or float(conf) < 0.35:
            continue
        xs = [p[0] for p in box]; ys = [p[1] for p in box]
        x0, x1, top, bot = min(xs) * k, max(xs) * k, min(ys) * k, max(ys) * k
        words.append({"text": _fix_ocr_spacing(txt.strip()), "x0": x0, "x1": x1, "top": top,
                      "bottom": bot, "size": max(6.0, (bot - top) * 0.8), "ocr": True})
    return words


LAST_OCR = False


def pdf_text(path):
    """Text of a PDF in reading order, columns separated, wide gaps marked."""
    import pdfplumber
    global LAST_OCR
    LAST_OCR = False
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            try:
                words = page.extract_words(keep_blank_chars=False, use_text_flow=False,
                                           extra_attrs=["size"])
            except Exception:
                words = []
            if len(words) < 8 and (page.images or not words):
                ocr = ocr_words(page)
                if ocr:
                    LAST_OCR = True
                words = ocr or words          # scanned page: no text layer
            if not words:
                t = page.extract_text() or ""
                if t:
                    pages.append(t)
                continue
            lines = _group_lines(words)
            gx = _find_gutter(lines, float(page.width))
            if gx is not None:
                lines = _split_columns(lines, gx)
            else:
                lines = _split_runs_of_two_blocks(lines)
            pages.append("\n".join(_line_text(ln) for ln in lines))
    return "\n".join(pages)


def docx_text(path):
    """Paragraphs and table rows of a .docx in document order."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []

    def walk(parent):
        for child in parent.element.body.iterchildren() if hasattr(parent.element, "body") else parent._element.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                out.append(Paragraph(child, parent).text)
            elif tag == "tbl":
                tbl = Table(child, parent)
                for row in tbl.rows:
                    cells, seen = [], set()
                    for c in row.cells:
                        if id(c._tc) in seen:
                            continue
                        seen.add(id(c._tc))
                        t = "\n".join(p.text for p in c.paragraphs if p.text.strip()).strip()
                        if t:
                            cells.append(t)
                    # one cell per line keeps cell content together
                    for c in cells:
                        out.extend(c.splitlines())
    try:
        walk(d)
    except Exception:
        out = [p.text for p in d.paragraphs]
    return "\n".join(out)


# ══════════════════════════════════════════════════════════════════════════════
# 2. Cleaning
# ══════════════════════════════════════════════════════════════════════════════

_CID_RE = re.compile(r"\(cid:\d+\)")
_BULLET_CHARS = "•●○◦▪■□▸►‣⁃∙·*"
_DASHES = "–—−"

_MON = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
_ONE = rf"(?:\d{{1,2}}/(?:19|20)\d{{2}}|(?:19|20)\d{{2}}-\d{{2}}|{_MON}\.?,?\s*(?:19|20)\d{{2}}|(?:19|20)\d{{2}})"
_END = r"(?:Present|Current|Now|Ongoing|Pursuing|Till\s+date|To\s+date|Today)"
DATE_RANGE = re.compile(rf"(?i)\b({_ONE})\s*(?:-|to|until|till)\s*({_ONE}|{_END})(?![\w/])")
DATE_SINGLE = re.compile(rf"(?i)\b({_ONE})(?![\w/])")
_DATE_ONLY = re.compile(rf"(?i)^\s*(?:{_ONE})(?:\s*(?:-|to)\s*(?:{_ONE}|{_END}))?\s*$")


def prepare_text(text):
    """Remove extraction artefacts and normalise punctuation."""
    if not text:
        return ""
    t = _CID_RE.sub("•", text)
    t = t.replace(" ", " ").replace("​", "").replace("", "•").replace("", "•")
    for d in _DASHES:
        t = t.replace(d, "-")
    t = re.sub(r"[ \t]+$", "", t, flags=re.M)

    # A date range split over two lines by a two-column layout:
    #   "2015-06 - Microsoft"          ->  "2015-06 - 2015-09 Microsoft"
    #   "2015-09 Software Engineer"        "Software Engineer"
    lines = t.splitlines()
    out, i = [], 0
    start_re = re.compile(rf"(?i)^(\s*)({_ONE})\s*-\s*(?P<rest>(?!{_ONE}|{_END}).*\S.*)?$")
    end_re = re.compile(rf"(?i)^\s*({_ONE}|{_END})(?:\s+(?P<rest>.*\S.*))?$")
    while i < len(lines):
        m = start_re.match(lines[i])
        if m and i + 1 < len(lines):
            n = end_re.match(lines[i + 1])
            if n and not DATE_RANGE.search(lines[i]):
                d1, d2 = m.group(2), n.group(1)
                r1 = (m.group("rest") or "").strip()
                r2 = (n.group("rest") or "").strip()
                out.append(f"{d1} - {d2} {r1}".rstrip())
                if r2:
                    out.append(r2)
                i += 2
                continue
        out.append(lines[i]); i += 1
    return "\n".join(out)


def _flat(s):
    return re.sub(r"\s+", " ", s).strip()


def _is_bullet(line):
    return bool(line.strip()) and line.strip()[0] in _BULLET_CHARS + "-"


def _strip_bullet(line):
    return re.sub(rf"^\s*[{re.escape(_BULLET_CHARS)}\-]+\s*", "", line).strip()


# ══════════════════════════════════════════════════════════════════════════════
# 3. Lexicons
# ══════════════════════════════════════════════════════════════════════════════

TITLE_WORDS = {
    "manager", "engineer", "developer", "intern", "internship", "trainee", "clerk", "specialist",
    "supervisor", "analyst", "technician", "phlebotomist", "consultant", "assistant", "associate",
    "executive", "officer", "director", "lead", "head", "architect", "administrator", "coordinator",
    "scientist", "designer", "programmer", "tester", "representative", "agent", "teacher", "lecturer",
    "professor", "researcher", "accountant", "auditor", "nurse", "driver", "operator", "technologist",
    "strategist", "planner", "recruiter", "advisor", "adviser", "founder", "owner", "freelancer",
    "apprentice", "fellow", "member", "volunteer", "mentor", "tutor", "writer", "editor", "cashier",
    "receptionist", "secretary", "pharmacist", "therapist", "paramedic", "foreman", "inspector",
    "controller", "bartender", "chef", "cook", "mechanic", "electrician", "sales", "cto", "ceo", "cfo",
    "vp", "president", "partner", "contractor", "scholar", "student", "ambassador", "generalist",
}
_TITLE_RE = re.compile(r"(?i)\b(" + "|".join(sorted(TITLE_WORDS, key=len, reverse=True)) + r")s?\b")

_ORG_SUFFIX = re.compile(
    r"(?i)\b(inc|llc|ltd|limited|pvt|private|corp|corporation|co|company|technologies|technology|"
    r"solutions|systems|labs|group|industries|services|consulting|consultants|bank|health|healthcare|"
    r"hospital|clinic|foundation|institute|university|college|school|academy|polytechnic|board|"
    r"association|society|council|agency|enterprises|software|networks|innovations|partners|"
    r"logistics|motors|airlines|airways|pharma|labs|studio|studios|ventures|capital|media)\b")

_INST_RE = re.compile(
    r"(?i)\b(university|college|institute|institution|school|academy|polytechnic|board|iit|nit|iiit|"
    r"vidyalaya|vidyapeeth|mahavidyalaya|campus|gymnasium|lyceum|conservatory|seminary)\b")
_GENERIC_EDU = {"high school", "secondary school", "senior secondary", "intermediate", "higher secondary",
                "school", "college", "university", "diploma", "graduation"}

_DEGREE_LINE = re.compile(
    r"(?i)(?<![A-Za-z0-9\.])(b\.?\s?tech|b\.?\s?e\.?|b\.?\s?sc?\.?|m\.?\s?tech|m\.?\s?sc?\.?|mba|ph\.?\s?d\.?|"
    r"b\.?\s?a\.?|m\.?\s?a\.?|bca|mca|b\.?com|m\.?com|bachelor|master|doctorate|associate|diploma|"
    r"certificate|pcm|pcb|hsc|ssc|cbse|icse|a-levels?|gcse|high\s*school|secondary|intermediate|"
    r"engineering|degree)(?![A-Za-z])")

_US_STATES = {
    "alabama", "alaska", "arizona", "arkansas", "california", "colorado", "connecticut", "delaware",
    "florida", "georgia", "hawaii", "idaho", "illinois", "indiana", "iowa", "kansas", "kentucky",
    "louisiana", "maine", "maryland", "massachusetts", "michigan", "minnesota", "mississippi",
    "missouri", "montana", "nebraska", "nevada", "new hampshire", "new jersey", "new mexico",
    "new york", "north carolina", "north dakota", "ohio", "oklahoma", "oregon", "pennsylvania",
    "rhode island", "south carolina", "south dakota", "tennessee", "texas", "utah", "vermont",
    "virginia", "washington", "west virginia", "wisconsin", "wyoming"}
_IN_STATES = {
    "andhra pradesh", "arunachal pradesh", "assam", "bihar", "chhattisgarh", "goa", "gujarat",
    "haryana", "himachal pradesh", "jharkhand", "karnataka", "kerala", "madhya pradesh",
    "maharashtra", "manipur", "meghalaya", "mizoram", "nagaland", "odisha", "punjab", "rajasthan",
    "sikkim", "tamil nadu", "telangana", "tripura", "uttar pradesh", "uttarakhand", "west bengal",
    "delhi", "chandigarh", "puducherry", "jammu and kashmir", "ladakh"}
_COUNTRIES = {
    "india", "usa", "united states", "u.s.a", "uk", "united kingdom", "england", "scotland", "wales",
    "canada", "australia", "germany", "france", "spain", "italy", "netherlands", "ireland", "singapore",
    "uae", "united arab emirates", "dubai", "china", "japan", "south korea", "brazil", "mexico",
    "south africa", "nigeria", "kenya", "egypt", "pakistan", "bangladesh", "sri lanka", "nepal",
    "philippines", "indonesia", "malaysia", "thailand", "vietnam", "new zealand", "sweden", "norway",
    "denmark", "finland", "switzerland", "austria", "belgium", "poland", "portugal", "turkey",
    "saudi arabia", "qatar", "israel", "russia", "ukraine", "argentina", "chile", "colombia"}
_REGIONS = _US_STATES | _IN_STATES | _COUNTRIES
_US_ABBR = ("AL AK AZ AR CA CO CT DE FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM "
            "NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY DC").split()

_CITY = r"[A-Z][A-Za-z\.'\-]+(?:\s[A-Z][A-Za-z\.'\-]+){0,2}"
_LOC_ABBR_RE = re.compile(rf"\b({_CITY}),\s*({'|'.join(_US_ABBR)})\b(?:,?\s*\d{{5,6}})?")
_LOC_REGION_RE = re.compile(
    rf"\b({_CITY}),\s*(" + "|".join(sorted((re.escape(r) for r in _REGIONS), key=len, reverse=True)) + r")\b", re.I)
_LOC_PIN_RE = re.compile(rf"\b({_CITY})\s+(\d{{6}})\b")
_LOC_PREFIX_RE = re.compile(
    rf"^\s*({_CITY},\s*(?:{'|'.join(_US_ABBR)})\b(?:,?\s*\d{{5}})?|{_CITY},\s*[A-Z][a-z]+(?:\s[A-Z][a-z]+)?)(?=\s|$)")

_SPOKEN = {
    "english", "hindi", "spanish", "french", "german", "italian", "portuguese", "russian", "chinese",
    "mandarin", "cantonese", "japanese", "korean", "arabic", "urdu", "bengali", "tamil", "telugu",
    "kannada", "malayalam", "marathi", "gujarati", "punjabi", "odia", "assamese", "nepali", "dutch",
    "swedish", "turkish", "polish", "greek", "hebrew", "persian", "farsi", "thai", "vietnamese",
    "indonesian", "malay", "swahili", "sanskrit", "tagalog", "konkani", "sindhi"}


# ══════════════════════════════════════════════════════════════════════════════
# 4. Line classification
# ══════════════════════════════════════════════════════════════════════════════

def has_title_word(s):
    return bool(_TITLE_RE.search(s))


def split_location_prefix(line):
    """'San Diego, CA Amazon' -> ('San Diego, CA', 'Amazon')."""
    line = _flat(line)
    m = _LOC_PREFIX_RE.match(line)
    if not m:
        return None, line
    loc = m.group(1)
    city_region = loc.split(",", 1)
    ok = False
    if len(city_region) == 2:
        reg = re.sub(r"[\s,\d]+$", "", city_region[1].strip()).strip()
        first = reg.split(",")[0].strip()
        ok = (first.upper() in _US_ABBR) or (first.lower() in _REGIONS)
    if not ok:
        return None, line
    return loc, line[m.end():].strip(" ,-|")


def _is_location_only(line):
    loc, rest = split_location_prefix(line)
    return bool(loc) and not rest


def looks_like_header_line(line):
    """A short, capitalised, non-sentence line that can be a title/company/degree."""
    s = _flat(_strip_bullet(line))
    if not s or len(s) > 90 or _is_bullet(line):
        return False
    if s[0].islower():
        return False
    if re.search(r"[.!?]$", s) and len(s.split()) > 7:
        return False
    if len(s.split()) > 11:
        return False
    return True


def _parse_dates(line):
    """Return (range_text, start, end, line_without_dates)."""
    m = DATE_RANGE.search(line)
    if m:
        return m.group(0), m.group(1), m.group(2), (line[:m.start()] + " " + line[m.end():])
    m = DATE_SINGLE.search(line)
    if m:
        return m.group(0), m.group(1), None, (line[:m.start()] + " " + line[m.end():])
    return None, None, None, line


def _is_anchor(line):
    """A line that carries the dates of an entry: it starts with, or consists of, a date."""
    s = line.strip()
    if not s or _is_bullet(s):
        return False
    rm = DATE_RANGE.search(s)
    if rm and len(s.split()) <= 16 and not re.search(r"[.!?]$", s):
        return True                       # "Role, Company  Jun 2023 - Aug 2023" (date anywhere)
    m = rm or DATE_SINGLE.search(s)
    if not m:
        return False
    if m.start() <= 2:                    # date first
        return True
    rest = (s[:m.start()] + s[m.end():]).strip(" -|,()")
    return len(rest.split()) <= 1 and not re.search(r"[A-Za-z]{4,}", rest)  # date-only (+ tiny label)


# ══════════════════════════════════════════════════════════════════════════════
# 5. Entry parsing
# ══════════════════════════════════════════════════════════════════════════════

def _entries(text, max_back=2):
    """Split a section into entries using date lines as anchors.

    Two common layouts are handled:
      forward  : "2016 - 2019  Certified Phlebotomist" / "Aurora Health" / bullets
      backward : "CRM Trainee" / "BetaEdge Innovations" / "Dec 2021 - Jan 2022" / bullets
    Returns a list of dicts: {dates, start, end, head:[...], body:[...]}.
    """
    lines = [l for l in text.splitlines()]
    idx = [i for i, l in enumerate(lines) if _is_anchor(l)]
    if not idx:
        return []
    entries = []
    heads = []          # (head_start, anchor_idx, header lines)
    for k, a in enumerate(idx):
        rng, s, e, rest = _parse_dates(lines[a])
        rest = _flat(rest).strip(" -|,()/·")
        loc_a = None
        if rest and _is_location_only(rest):
            loc_a, rest = rest, ""
        header, h0 = [], a
        if rest:                                   # forward style
            header = [rest]
            j = a + 1
            while j < len(lines) and len(header) < 3 and (k + 1 >= len(idx) or j < idx[k + 1]):
                L = lines[j]
                if not L.strip():
                    j += 1; continue
                if _is_bullet(L) or not looks_like_header_line(L) or re.search(r"[.!?]$", L.strip()) \
                        or re.match(r"(?i)^(built|led|ran|wrote|made|worked|\w+ed|\w+ing)\b", L.strip()):
                    break
                header.append(_flat(L)); j += 1
        else:                                      # backward style
            j = a - 1
            prev_anchor = idx[k - 1] if k else -1
            tmp = []
            while j > prev_anchor and len(tmp) < max_back:
                L = lines[j]
                if not L.strip():
                    j -= 1
                    if tmp: break
                    continue
                if _is_bullet(L) or _is_anchor(L) or re.match(r"(?i)^(built|led|ran|wrote|made|worked|\w+ed|\w+ing)\b", L.strip()):
                    break
                if not looks_like_header_line(L):
                    if tmp or not (len(L.split()) <= 10 and not re.search(r"[.!?]$", L.strip())):
                        break
                tmp.insert(0, _flat(L)); h0 = j; j -= 1
            header = tmp
        heads.append((h0, a, header, rng, s, e, loc_a))
    for k, (h0, a, header, rng, s, e, loc_a) in enumerate(heads):
        forward = bool(header) and bool(_flat(_parse_dates(lines[a])[3]).strip(" -|,()/·")) and not loc_a
        body_start = a + 1 + (max(0, len(header) - 1) if forward else 0)
        # body ends where the next entry's header begins
        body_end = heads[k + 1][0] if k + 1 < len(heads) else len(lines)
        body = [l for l in lines[body_start:body_end] if l.strip()]
        entries.append({"dates": rng, "start": s, "end": e, "head": header, "body": body, "loc": loc_a})
    if entries:
        entries[0]["lead_lines"] = [l for l in lines[:heads[0][0]] if l.strip()]
    return entries


def _bullets_of(body):
    """Merge wrapped lines into bullet strings."""
    out = []
    for l in body:
        if _is_bullet(l):
            out.append(_strip_bullet(l))
        elif out and (l.strip()[:1].islower() or not re.search(r"[.!?]$", out[-1])):
            out[-1] += ("" if out[-1].endswith("-") and l.strip()[:1].islower() else " ") + l.strip()
        else:
            out.append(l.strip())
    return [_flat(b) for b in out if b.strip()]


def parse_experience(text):
    """-> [{title, company, location, dates, start, end, bullets, description, years, dates_mentioned}]"""
    res = []
    for e in _entries(text):
        title = company = location = None
        cands = []
        for h in e["head"]:
            loc, rest = split_location_prefix(h)
            if loc and not location:
                location = loc
            if rest:
                cands.append(rest)
        # "Role at Company" / "Role, Company" / "Role | Company"
        if len(cands) == 1:
            m = re.match(r"^(.+?)\s+(?:at|@|\||-)\s+(.+)$", cands[0]) or re.match(r"^(.+?),\s+(.+)$", cands[0])
            if m and (has_title_word(m.group(1)) or has_title_word(m.group(2))):
                cands = [m.group(1).strip(), m.group(2).strip()]
        for i, c in enumerate(cands):
            if "," in c and not has_title_word(c.split(",")[0]):
                first, rest = c.split(",", 1)
                cands[i] = first.strip()
                location = location or rest.strip(" ,")
        t_idx = next((i for i, c in enumerate(cands) if has_title_word(c)), None)
        if t_idx is not None:
            title = cands[t_idx]
            others = [c for i, c in enumerate(cands) if i != t_idx]
            company = others[0] if others else None
        elif cands:
            title = cands[0]
            company = cands[1] if len(cands) > 1 else None
        if company and has_title_word(company) and not _ORG_SUFFIX.search(company):
            pass
        bullets = _bullets_of(e["body"])
        body_loc = None
        res.append({
            "title": title, "company": company, "location": location or e.get("loc") or body_loc,
            "dates": e["dates"], "start": e["start"], "end": e["end"],
            "bullets": bullets,
            "description": _flat(" ".join(filter(None, [title, company, e["dates"]])) + ". " + " ".join(bullets))[:600],
            "years": re.findall(r"(?:19|20)\d{2}", e["dates"] or ""),
            "dates_mentioned": [e["dates"]] if e["dates"] else [],
        })
    return res


_ABBR_DEGREE = re.compile(r"(?i)^(?:[BM]\.?\s?[A-Za-z]{1,4}\.?|bachelors?|masters?|ph\.?d\.?|doctorate|diploma|associate)$")


def parse_education(text):
    """-> [{degree, field, institution, location, years, dates, gpa}]"""
    res = []
    for e in _entries(text, max_back=3):
        degree = inst = location = e.get("loc")
        degree = inst = None
        heads = list(e["head"])
        # "B.S." on one line and "Computer Science" on the next -> "B.S. Computer Science"
        merged = []
        for h in heads:
            if merged and _ABBR_DEGREE.match(merged[-1].strip()) and not _INST_RE.search(h) and not _is_location_only(h):
                merged[-1] = merged[-1].strip() + " " + h
            else:
                merged.append(h)
        for i in range(len(merged) - 1):
            if _INST_RE.search(merged[i]) and _ABBR_DEGREE.match(merged[i + 1].strip()):
                pass
        heads = merged
        for h in heads:
            loc, rest = split_location_prefix(h)
            if loc and not location:
                location = loc
            h = rest
            if not h:
                continue
            # "University of Oregon, M.S. Computer Science"
            parts = [p.strip() for p in re.split(r",\s*|\s+-\s+|\s+\|\s+", h) if p.strip()]
            for p in parts:
                is_inst = bool(_INST_RE.search(p)) and p.lower() not in _GENERIC_EDU
                is_deg = bool(_DEGREE_LINE.search(p)) and not (is_inst and not re.search(r"(?i)diploma|bachelor|master|b\.?\s?tech|m\.?\s?tech", p))
                if p.lower() in _GENERIC_EDU and inst is None and degree is None and p.lower() in {"high school", "secondary school", "senior secondary", "higher secondary", "intermediate"}:
                    degree = degree or p
                    continue
                if is_inst and inst is None:
                    inst = p
                elif is_deg and degree is None:
                    degree = p
                elif inst is None and degree is not None and not is_deg:
                    inst = p
                elif degree is None and inst is not None and not is_inst:
                    degree = p
                elif inst is None and degree is None:
                    degree = p
        for l in e["body"][:3]:
            if _is_location_only(l) and not location:
                location = _flat(l)
        # an institution line like "Winchester East High School" is the entry title when no degree line
        body_txt = " ".join(e["body"])
        gpa = None
        m = re.search(r"(?i)\b(?:gpa|cgpa|sgpa|percentage|score|grade)\s*[:\-]?\s*([\d\.]+\s*(?:/\s*\d+(?:\.\d+)?|%)?)", body_txt)
        if m:
            gpa = m.group(1).strip()
        # degree sometimes lives in the body: "High School Diploma"
        if not degree:
            for l in e["body"]:
                if _is_bullet(l):
                    continue
                if _DEGREE_LINE.search(l) and len(l.split()) <= 10:
                    degree = _flat(l); break
        elif inst and degree and degree.lower() == inst.lower():
            degree = None
        # for an entry with a high-school institution and a "Diploma" body line
        if inst and degree is None:
            for l in e["body"]:
                if re.search(r"(?i)diploma|certificate", l) and len(l.split()) <= 6:
                    degree = _flat(l); break
        if not (degree or inst):
            continue
        full = degree or inst
        field = None
        m = re.search(r"\bin\s+([A-Z][A-Za-z&,\s]+)$", degree or "") or \
            re.search(r"(?i)^(?:b\.?\s?tech|m\.?\s?tech|b\.?e\.?|m\.?s\.?|b\.?s\.?|diploma|ph\.?d\.?|mba|bca|mca)\.?\s+(?:in\s+)?([A-Z][A-Za-z&,\s]+)$", degree or "")
        if m:
            field = m.group(1).strip()
        res.append({
            "degree": full if not (degree and inst) else degree,
            "institution": inst, "field": field, "location": location or e.get("loc"),
            "years": re.findall(r"(?:19|20)\d{2}", e["dates"] or ""),
            "dates": e["dates"], "gpa": gpa,
        })
    return res


def parse_titled_items(text):
    """Certificates / courses / projects written as  Title / description / date
    (or date-less bullets). -> [{title, description, dates}]"""
    items = []
    ents = _entries(text, max_back=3)
    if ents:
        lead = [l for l in ents[0].get("lead_lines", []) if l.strip()]
        k = 0
        while k < len(lead) and k < 2 and looks_like_header_line(lead[k]) and not _is_bullet(lead[k]):
            k += 1
        if k:
            items.append({"title": _flat(lead[0]), "dates": None, "subtitle": _flat(" ".join(lead[1:k])),
                          "description": _flat(" ".join(_bullets_of(lead[k:])))})
    for e in ents:
        head = [h for h in e["head"] if h]
        if not head:
            continue
        title = head[0]
        sub = " ".join(head[1:])
        items.append({"title": title, "subtitle": _flat(sub), "description": _flat(" ".join(_bullets_of(e["body"]))), "dates": e["dates"]})
    return items


def parse_certifications(text):
    items = parse_titled_items(text)
    if items:
        out = []
        for it in items:
            s = it["title"]
            if it.get("subtitle"):
                s += " - " + it["subtitle"]
            if it["description"]:
                s += " - " + it["description"]
            if it["dates"]:
                s += f" ({it['dates']})"
            out.append(s)
        return out
    out = []
    for l in text.splitlines():
        l = _strip_bullet(l)
        if len(l) >= 4:
            out.append(l)
    return out[:20]


def parse_achievements(text):
    """Cards ('Title' + description lines) or plain bullets."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return []
    if any(_is_bullet(l) for l in lines):
        return [_strip_bullet(l) for l in lines if len(_strip_bullet(l)) >= 4][:20]
    items = []
    for l in lines:
        cont = bool(items) and (len(l.split()) >= 8 or l[:1].islower() or re.search(r"[.!?]$", l))
        if cont:
            items[-1] += (" - " if " - " not in items[-1] and len(items[-1].split()) <= 7 else " ") + l
        else:
            items.append(l)
    return [_flat(i) for i in items][:20]


# ══════════════════════════════════════════════════════════════════════════════
# 6. Header, summary, skills, languages
# ══════════════════════════════════════════════════════════════════════════════

_CONTACT_RE = re.compile(r"(?i)@|https?://|www\.|linkedin|github|\+?\d[\d\s\-\.\(\)]{8,}\d|\b\d{6}\b")


def header_lines(text, n=18):
    """First lines of the resume; wide gaps are kept as double spaces."""
    out = []
    for l in text.splitlines()[:n]:
        l = re.sub(r"[ \t]{3,}", _GAP, l.strip())
        if l:
            out.append(re.sub(r"(?<! ) (?! )", " ", l))
    return out


def extract_headline(lines, name):
    """The line (or two) right under the name: 'Certified Phlebotomy Technician'."""
    if not name:
        return None
    lines = [_flat(l) for l in lines]
    for i, l in enumerate(lines):
        if name.lower() in l.lower() and len(l) <= len(name) + 4:
            nxt = lines[i + 1:i + 3]
            if not nxt:
                return None
            h = nxt[0]
            if _CONTACT_RE.search(h) or len(h.split()) > 14 or re.search(r"[.!?]$", h):
                return None
            if re.match(r"(?i)^(summary|objective|profile|experience|education|skills|address|phone|e-?mail)\b", h):
                return None
            if len(nxt) > 1 and len(nxt[1].split()) <= 3 and nxt[1][:1].isupper() \
                    and not _CONTACT_RE.search(nxt[1]) and not re.search(r"[.!?:,]$", nxt[1]) \
                    and not _is_section_header(nxt[1]) and not re.search(r"(?i)^(address|phone|e-?mail|www|linkedin|twitter)$", nxt[1]):
                h += " " + nxt[1]
            return h
    return None


def extract_header_location(lines):
    """City/region from the contact area (first lines of the resume)."""
    blob = [l for l in lines if not re.search(r"(?i)^\s*(summary|objective|experience|education)\b", l)]
    for l in blob:
        l2 = re.sub(r"(?i)^(address|location|city)\s*[:\-]?\s*", "", l)
        for seg in re.split(r"\s{2,}|\|", l2):
            seg = seg.strip()
            for rx in (_LOC_ABBR_RE, _LOC_REGION_RE):
                m = rx.search(seg)
                if m and not has_title_word(m.group(1)) and not _ORG_SUFFIX.search(m.group(1)):
                    return _flat(re.sub(r"(?i)^(address|location)\s*[:\-]?\s*", "", m.group(0)))
            seg = re.sub(r"^[A-Z](?=[A-Z][a-z])", "", seg)
            m = _LOC_PIN_RE.search(seg)
            if m and not has_title_word(m.group(1)):
                return _flat(m.group(0))
    return None


def extract_summary_fallback(other_text, name, headline):
    """Opening paragraph without a 'Summary' heading."""
    lines = []
    for l in other_text.splitlines():
        s = _flat(l)
        if not s or _CONTACT_RE.search(s) or (name and s.lower() == name.lower()):
            continue
        if headline and s in headline:
            continue
        lines.append(s)
    para = [s for s in lines if len(s.split()) >= 7]
    return " ".join(para)[:900] if len(" ".join(para).split()) >= 12 else ""


_HEAD_RE = re.compile(
    r"(?i)^\W*(summary|objective|profile|about\s*me|experience|work\s*experience|education|skills?|"
    r"technical\s*skills?|projects?|certificat\w*|courses?|achievements?|key\s*achievements?|languages?|interests?|"
    r"hobbies|references?|awards?|publications?|volunteer\w*|activities|contact|contact\s*(details|information)|personal\s*(details|information))\W*$")


def _is_section_header(line):
    return bool(_HEAD_RE.match(line.strip()))


_SKILL_NOISE = {"and", "or", "etc", "others", "other", "skills", "tools", "more", "various"}


_JOINED = {"sharepoint", "workflows", "powerapps", "wordpress", "javascript", "typescript", "github", "gitlab",
           "powerbi", "nodejs", "mongodb", "postgresql", "mysql", "photoshop", "bootstrap", "jquery", "devops"}


def parse_listed_skills(text):
    """Every skill-like item written in a skills section, in the resume's own words
    ('Patient Rapport Building', 'Business Rules'). Not limited to the vocabulary."""
    if not text.strip():
        return []
    sub_re = re.compile(r"(?i)^\W*(hard|soft|technical|core|other|key|programming|professional|personal|interpersonal)\s*(skills?|competenc\w+)?\W*$")
    lines = [l for l in text.splitlines() if l.strip() and not sub_re.match(l.strip())]
    merged_lines = []
    for l in lines:
        if merged_lines and (merged_lines[-1].rstrip().endswith("-") or
                             (re.fullmatch(r"[A-Za-z]+-[A-Za-z]+", merged_lines[-1].strip()) and l.strip()[:1].isupper() and len(l.split()) <= 4)):
            glue = "" if merged_lines[-1].rstrip().endswith("-") else " "
            merged_lines[-1] = merged_lines[-1].rstrip() + glue + l.strip()
        else:
            merged_lines.append(l)
    lines = merged_lines
    comma_mode = any(("," in re.sub(r"\([^)]*\)", "", l)) or ";" in l for l in lines)
    chunks = []
    if comma_mode:
        # soft-wrapped comma lists continue on the next line; separate rows do not
        longest = max(len(l) for l in lines)
        try:
            from src.resume_parser import SKILLS_DB
            known = {k.lower() for k in SKILLS_DB}
        except Exception:
            known = set()
        groups, cur, n_cur = [], "", 0
        for l in lines:
            l = _strip_bullet(l) if _is_bullet(l) else l.strip()
            last_w = re.sub(r"\W", "", cur.split()[-1]).lower() if cur.split() else ""
            first_w = re.sub(r"\W", "", l.split()[0]).lower() if l.split() else ""
            both_known = last_w in known and first_w in known
            if cur and (cur.rstrip().endswith(",") or (n_cur == 1 and not both_known and len(cur) >= 0.85 * longest and len(lines) > 1 and not cur.rstrip().endswith("."))):
                cur += " " + l; n_cur += 1
            else:
                if cur: groups.append(cur)
                cur, n_cur = l, 1
        if cur: groups.append(cur)
        blob = "  ".join(groups)
        depth, cur = 0, ""
        for ch in blob:
            depth += (ch == "(") - (ch == ")" and depth > 0)
            if depth == 0 and ch in ",;|•":
                chunks.append(cur); cur = ""
            else:
                cur += ch
        chunks.append(cur)
        chunks = [p for c in chunks for p in re.split(r"\s{2,}", c)]
    else:
        for l in lines:
            l = _strip_bullet(l) if _is_bullet(l) else l
            chunks.extend(re.split(r"\s{2,}|[|•;]|(?<=\b[A-Z])\.\s+(?=[A-Z])", l))
    try:
        from src.resume_parser import SKILLS_DB
        known_db = {k.lower() for k in SKILLS_DB}
    except Exception:
        known_db = set()
    out, seen = [], set()
    for c in chunks:
        c = _flat(re.sub(r"^[^\w(]+|[^\w)]+$", "", c))
        if c.count("(") != c.count(")"):
            c = c.replace("(", "").replace(")", "")
        c = re.sub(r"^(?:and|&)\s+", "", c, flags=re.I)
        parts = c.split()
        if len(parts) == 2 and (parts[0] + parts[1]).lower() in (known_db | _JOINED) and True:
            c = parts[0] + parts[1]              # Share Point -> SharePoint, Work flows -> Workflows
        low = c.lower()
        if not c or len(c) > 42 or len(c.split()) > 5 or low in _SKILL_NOISE or low in _SPOKEN:
            continue
        if re.search(r"[.!?]$", c) or re.fullmatch(r"[\d\W]+", c):
            continue
        if low in seen:
            continue
        seen.add(low); out.append(c)
    return out[:60]


def parse_languages(text):
    found = []
    for tok in re.split(r"[,;|•\n]|\s{2,}", text or ""):
        t = re.sub(r"(?i)\b(native|fluent|conversational|basic|intermediate|advanced|proficient|beginner|bilingual)\b", "", tok)
        t = _flat(re.sub(r"[\(\)\-:]", " ", t))
        for w in re.findall(r"[A-Za-z]+", t):
            if w.lower() in _SPOKEN and w.title() not in found:
                found.append(w.title())
    return found


# ══════════════════════════════════════════════════════════════════════════════
# 7. Entities derived from the structure
# ══════════════════════════════════════════════════════════════════════════════

def clean_org(s):
    s = _flat(s or "")
    s = re.sub(r"\s*\(.*?\)\s*$", "", s).strip(" ,-|")
    return s


def issuer_orgs(cert_lines):
    """'First Aid, CPR—American Red Cross' -> American Red Cross; '... (National Healthcareer Association)'"""
    out = []
    for c in cert_lines:
        for m in re.finditer(r"\(([A-Z][^()]{3,60})\)", c):
            out.append(m.group(1).strip())
        m = re.search(r"(?:-|by|from|\bat\b)\s+((?:[A-Z][\w&\.\-']+\s?){2,6})\s*(?:\(|$|,)", c)
        if m:
            out.append(m.group(1).strip())
    res, seen = [], set()
    for o in out:
        o = clean_org(o)
        if o and not DATE_SINGLE.search(o) and not re.search(r"\d", o) and o.lower() not in seen and len(o.split()) <= 6:
            seen.add(o.lower()); res.append(o)
    return res
